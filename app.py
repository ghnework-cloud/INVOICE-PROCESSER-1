# VERSION: 20261018_varsha_excel
import streamlit as st
import tempfile, os, sys
from pathlib import Path

sys.path.insert(0, os.path.dirname(__file__))

st.set_page_config(
    page_title="Invoice Processor",
    page_icon="▣",
    layout="centered",
    initial_sidebar_state="collapsed"
)

# ── Custom CSS ──────────────────────────────────────────────────────────────
st.markdown("""
<style>
@import url('https://fonts.googleapis.com/css2?family=IBM+Plex+Mono:wght@400;700&family=IBM+Plex+Sans:wght@400;600;700&display=swap');

html, body, [class*="css"] {
    font-family: 'IBM Plex Sans', sans-serif;
    background-color: #0a0a0a;
    color: #f0f0f0;
}

.main { background: #0a0a0a; }
.block-container { padding-top: 2rem; max-width: 720px; }

/* Header */
.app-header {
    display: flex; align-items: center; gap: 14px;
    padding-bottom: 20px;
    border-bottom: 1px solid #1e1e1e;
    margin-bottom: 28px;
}
.logo {
    width: 44px; height: 44px;
    background: #F5C518;
    border-radius: 8px;
    display: flex; align-items: center; justify-content: center;
    font-size: 22px; font-weight: 700; color: #000;
}
.app-title { font-size: 22px; font-weight: 700; color: #fff; }
.app-sub { font-size: 13px; color: #555; margin-top: 2px; }

/* Upload area */
[data-testid="stFileUploader"] {
    background: #111 !important;
    border: 2px dashed #2a2a2a !important;
    border-radius: 12px !important;
    padding: 8px !important;
}
[data-testid="stFileUploader"]:hover {
    border-color: #F5C518 !important;
}

/* Process button */
.stButton > button {
    background: #F5C518 !important;
    color: #000 !important;
    font-weight: 700 !important;
    font-size: 16px !important;
    border: none !important;
    border-radius: 10px !important;
    padding: 14px 24px !important;
    width: 100% !important;
    letter-spacing: 0.5px;
    transition: all 0.2s;
}
.stButton > button:hover {
    background: #FFD740 !important;
    transform: translateY(-1px);
}

/* Download button */
[data-testid="stDownloadButton"] > button {
    background: #22C55E !important;
    color: #000 !important;
    font-weight: 700 !important;
    font-size: 15px !important;
    border: none !important;
    border-radius: 10px !important;
    padding: 12px 24px !important;
    width: 100% !important;
}

/* Log area */
.log-box {
    background: #080808;
    border: 1px solid #1a1a1a;
    border-radius: 8px;
    padding: 14px 16px;
    font-family: 'IBM Plex Mono', monospace;
    font-size: 12px;
    line-height: 1.7;
    max-height: 300px;
    overflow-y: auto;
    white-space: pre;
    color: #ccc;
}

/* Status badges */
.badge-ok   { color: #22C55E; }
.badge-err  { color: #EF4444; }
.badge-warn { color: #F5C518; }
.badge-grey { color: #555; }

/* Cards */
.info-card {
    background: #111;
    border: 1px solid #1e1e1e;
    border-radius: 10px;
    padding: 16px 20px;
    margin-bottom: 12px;
}
.stat { font-size: 28px; font-weight: 700; color: #F5C518; }
.stat-label { font-size: 12px; color: #555; margin-top: 2px; }

/* Hide streamlit branding */
#MainMenu, footer, header { visibility: hidden; }
</style>
""", unsafe_allow_html=True)

# ── Header ──────────────────────────────────────────────────────────────────
st.markdown("""
<div class="app-header">
  <div class="logo">▣</div>
  <div>
    <div class="app-title">Invoice Processor</div>
    <div class="app-sub">Amazon Seller — Auto-classify &amp; stamp invoices</div>
  </div>
</div>
""", unsafe_allow_html=True)

# ── Varsha Excel generator ───────────────────────────────────────────────────
def is_varsha_pdf(pdf_bytes):
    """Check if this PDF is from Varsha Hardikkumar Raval."""
    try:
        import io
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(pdf_bytes))
        for page in reader.pages[:10]:  # check first 10 pages
            text = page.extract_text() or ""
            if "VARSHA HARDIKKUMAR RAVAL" in text.upper():
                return True
    except Exception:
        pass
    return False


def build_varsha_excel(pdf_bytes, filename_stem):
    """
    Extract invoice date + product description + qty from Varsha's PDF
    and return an Excel workbook as bytes.
    """
    import io
    import re
    from pypdf import PdfReader
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from collections import defaultdict
    from datetime import datetime

    reader = PdfReader(io.BytesIO(pdf_bytes))
    rows = []  # list of (date_str, product_name, qty)

    i = 0
    pages = reader.pages
    while i < len(pages):
        text = pages[i].extract_text() or ""

        # Detect invoice page (has "Tax Invoice" and "Sold By")
        if "Tax Invoice" in text and ("VARSHA HARDIKKUMAR RAVAL" in text.upper() or "Sold By" in text):
            # Extract invoice date
            date_match = re.search(r'Invoice Date\s*[:\-]?\s*(\d{1,2}[./\-]\d{1,2}[./\-]\d{4})', text)
            if not date_match:
                date_match = re.search(r'Order Date\s*[:\-]?\s*(\d{1,2}[./\-]\d{1,2}[./\-]\d{4})', text)
            raw_date = date_match.group(1) if date_match else ""
            # Normalise to DD.MM.YYYY
            if raw_date:
                for fmt in ("%d.%m.%Y", "%d/%m/%Y", "%d-%m-%Y", "%m/%d/%Y"):
                    try:
                        dt = datetime.strptime(raw_date, fmt)
                        date_str = dt.strftime("%d.%m.%Y")
                        break
                    except ValueError:
                        date_str = raw_date
            else:
                date_str = ""

            # Qty: unit price qty net pattern e.g. ₹171.42 1 ₹171.42
            qty_match = re.search(r'₹[\d,]+\.?\d*\s+(\d+)\s+₹', text)
            qty = int(qty_match.group(1)) if qty_match else 1

            # Product name: after row number '1 ', before ASIN (B0xxx) or HSN:
            prod_match = re.search(r'\n1\s+(.+?)(?:\n[A-Z0-9]{10,}\s*\(|HSN:|₹)', text, re.DOTALL)
            if prod_match:
                raw = prod_match.group(1).strip()
                product_name = re.sub(r'\s+', ' ', raw).strip()
                # Trim ASIN/bracket parts after last meaningful pipe segment
                if '|' in product_name:
                    parts = product_name.split('|')
                    clean_parts = []
                    for p in parts:
                        p = p.strip()
                        if re.match(r'^B0[A-Z0-9]{8}', p) or re.match(r'^[A-Z0-9]{2}-[A-Z0-9]', p):
                            break
                        clean_parts.append(p)
                    product_name = ' | '.join(clean_parts).strip(' |')
            else:
                product_name = "Unknown"

            if date_str or product_name != "Unknown":
                rows.append((date_str, product_name, qty))
        i += 1

    # ── Build workbook ────────────────────────────────────────────────────
    wb = openpyxl.Workbook()

    # ── Sheet 1: Packing Log ──────────────────────────────────────────────
    ws1 = wb.active
    ws1.title = "Packing Log"

    header_fill = PatternFill("solid", fgColor="F5C518")
    header_font = Font(name="Arial", bold=True, size=11, color="000000")
    data_font   = Font(name="Arial", size=10)
    thin = Side(style="thin", color="CCCCCC")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    headers = ["Date", "Product Name", "Quantity"]
    col_widths = [14, 70, 12]

    for col, (h, w) in enumerate(zip(headers, col_widths), 1):
        cell = ws1.cell(row=1, column=col, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border
        ws1.column_dimensions[cell.column_letter].width = w

    ws1.row_dimensions[1].height = 22

    for r, (date_str, product_name, qty) in enumerate(rows, 2):
        ws1.cell(row=r, column=1, value=date_str).font = data_font
        ws1.cell(row=r, column=1).alignment = Alignment(horizontal="center", vertical="top")
        ws1.cell(row=r, column=1).border = border

        ws1.cell(row=r, column=2, value=product_name).font = data_font
        ws1.cell(row=r, column=2).alignment = Alignment(wrap_text=True, vertical="top")
        ws1.cell(row=r, column=2).border = border

        ws1.cell(row=r, column=3, value=qty).font = data_font
        ws1.cell(row=r, column=3).alignment = Alignment(horizontal="center", vertical="top")
        ws1.cell(row=r, column=3).border = border

    # ── Sheet 2: Monthly Summary ──────────────────────────────────────────
    ws2 = wb.create_sheet("Monthly Summary")

    summary_font = Font(name="Arial", bold=True, size=11, color="000000")
    total_fill   = PatternFill("solid", fgColor="222222")
    total_font   = Font(name="Arial", bold=True, size=11, color="F5C518")

    ws2.column_dimensions["A"].width = 70
    ws2.column_dimensions["B"].width = 22

    for col, h in enumerate(["Product Name", "Total Quantity Packed"], 1):
        cell = ws2.cell(row=1, column=col, value=h)
        cell.font = summary_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border
    ws2.row_dimensions[1].height = 22

    # Aggregate
    totals = defaultdict(int)
    for _, product_name, qty in rows:
        totals[product_name] += qty

    grand_total = 0
    for r, (prod, total) in enumerate(sorted(totals.items()), 2):
        ws2.cell(row=r, column=1, value=prod).font = data_font
        ws2.cell(row=r, column=1).alignment = Alignment(wrap_text=True, vertical="top")
        ws2.cell(row=r, column=1).border = border
        ws2.cell(row=r, column=2, value=total).font = data_font
        ws2.cell(row=r, column=2).alignment = Alignment(horizontal="center", vertical="top")
        ws2.cell(row=r, column=2).border = border
        grand_total += total

    # Grand Total row
    last_row = len(totals) + 2
    gt1 = ws2.cell(row=last_row, column=1, value="Grand Total")
    gt1.font = total_font
    gt1.fill = total_fill
    gt1.alignment = Alignment(horizontal="right", vertical="center")
    gt1.border = border
    gt2 = ws2.cell(row=last_row, column=2, value=grand_total)
    gt2.font = total_font
    gt2.fill = total_fill
    gt2.alignment = Alignment(horizontal="center", vertical="center")
    gt2.border = border

    # Save to bytes
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.read()


# ── Upload ──────────────────────────────────────────────────────────────────
uploaded = st.file_uploader(
    "Drop your Amazon invoice PDF here",
    type=["pdf"],
    help="Upload your Amazon seller invoice batch PDF"
)

if uploaded:
    size_kb = len(uploaded.getvalue()) / 1024
    size_str = f"{size_kb:.0f} KB" if size_kb < 1024 else f"{size_kb/1024:.1f} MB"
    st.markdown(f"""
    <div class="info-card" style="display:flex;align-items:center;gap:14px;">
      <div style="font-size:28px">📄</div>
      <div>
        <div style="font-weight:600;color:#fff">{uploaded.name}</div>
        <div style="font-size:12px;color:#555;margin-top:2px">{size_str}</div>
      </div>
      <div style="margin-left:auto;color:#22C55E;font-weight:700">✓ Ready</div>
    </div>
    """, unsafe_allow_html=True)

    if st.button("⚡   PROCESS INVOICES"):
        with st.spinner("Processing…"):
            try:
                import importlib, sys
                if 'processor' in sys.modules:
                    del sys.modules['processor']
                if '_proc_mod' in sys.modules:
                    del sys.modules['_proc_mod']
                from processor import parse_pdf, build_output_pdf, double_check

                pdf_bytes_in = uploaded.getvalue()

                # Save uploaded PDF to temp file
                with tempfile.NamedTemporaryFile(suffix=".pdf", delete=False) as tmp_in:
                    tmp_in.write(pdf_bytes_in)
                    in_path = tmp_in.name

                out_path = in_path.replace(".pdf", "_PROCESSED.pdf")

                # Process
                log_lines = []
                reader, groups = parse_pdf(in_path)
                log_lines.append(f"✓  Found {len(groups)} shipments ({len(reader.pages)} total pages)")

                for g in groups:
                    tag = "✓" if g["confidence"] == "HIGH" else "⚠"
                    log_lines.append(
                        f"   {g['inv_num']:12s} → {g['product']:30s} qty={g['qty']}"
                    )

                problems = double_check(groups)

                if problems:
                    st.warning(f"⚠️ {len(problems)} issue(s) found — review below (output still generated)")
                    for p in problems:
                        st.warning(f"  {p}")

                build_output_pdf(in_path, out_path, groups)

                with open(out_path, "rb") as f:
                    pdf_bytes = f.read()

                os.unlink(in_path)
                os.unlink(out_path)

                # Stats
                from pypdf import PdfReader
                import io
                pages = len(PdfReader(io.BytesIO(pdf_bytes)).pages)

                col1, col2, col3 = st.columns(3)
                with col1:
                    st.markdown(f"""
                    <div class="info-card" style="text-align:center">
                      <div class="stat">{len(groups)}</div>
                      <div class="stat-label">SHIPMENTS</div>
                    </div>""", unsafe_allow_html=True)
                with col2:
                    st.markdown(f"""
                    <div class="info-card" style="text-align:center">
                      <div class="stat">{pages}</div>
                      <div class="stat-label">PAGES</div>
                    </div>""", unsafe_allow_html=True)
                with col3:
                    issue_label = f"⚠️ {len(problems)}" if problems else "✓"
                    issue_color = "#F59E0B" if problems else "#22C55E"
                    st.markdown(f"""
                    <div class="info-card" style="text-align:center">
                      <div class="stat" style="color:{issue_color}">{issue_label}</div>
                      <div class="stat-label">{"WARNINGS" if problems else "ZERO ISSUES"}</div>
                    </div>""", unsafe_allow_html=True)

                st.success("✅  Processing complete!")

                out_name = Path(uploaded.name).stem + "_PROCESSED.pdf"

                # ── Varsha Excel ──────────────────────────────────────────
                varsha = is_varsha_pdf(pdf_bytes_in)
                if varsha:
                    col_pdf, col_xl = st.columns(2)
                    with col_pdf:
                        st.download_button(
                            label="⬇   DOWNLOAD PDF",
                            data=pdf_bytes,
                            file_name=out_name,
                            mime="application/pdf"
                        )
                    with col_xl:
                        try:
                            xl_bytes = build_varsha_excel(pdf_bytes_in, Path(uploaded.name).stem)
                            xl_name = Path(uploaded.name).stem + "_Packing_Log.xlsx"
                            st.download_button(
                                label="📊   DOWNLOAD EXCEL",
                                data=xl_bytes,
                                file_name=xl_name,
                                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
                            )
                        except Exception as xe:
                            st.warning(f"Excel generation failed: {xe}")
                else:
                    st.download_button(
                        label="⬇   DOWNLOAD PROCESSED PDF",
                        data=pdf_bytes,
                        file_name=out_name,
                        mime="application/pdf"
                    )

                # Show log
                st.markdown("**Log:**")
                log_text = "\n".join(log_lines)
                st.markdown(f'<div class="log-box">{log_text}</div>',
                            unsafe_allow_html=True)

            except Exception as e:
                import traceback
                st.error(f"Error: {e}")
                st.code(traceback.format_exc())

else:
    # Info cards when no file selected
    st.markdown("""
    <div style="display:grid;grid-template-columns:1fr 1fr 1fr;gap:12px;margin-top:8px">
      <div class="info-card" style="text-align:center">
        <div style="font-size:28px;margin-bottom:8px">🏷</div>
        <div style="font-weight:600;color:#fff;font-size:13px">Auto Classify</div>
        <div style="font-size:11px;color:#555;margin-top:4px">150+ product codes</div>
      </div>
      <div class="info-card" style="text-align:center">
        <div style="font-size:28px;margin-bottom:8px">✅</div>
        <div style="font-weight:600;color:#fff;font-size:13px">Zero Mistakes</div>
        <div style="font-size:11px;color:#555;margin-top:4px">Double-check validator</div>
      </div>
      <div class="info-card" style="text-align:center">
        <div style="font-size:28px;margin-bottom:8px">⭐</div>
        <div style="font-weight:600;color:#fff;font-size:13px">Multi-Qty Stars</div>
        <div style="font-size:11px;color:#555;margin-top:4px">Staff alert system</div>
      </div>
    </div>
    """, unsafe_allow_html=True)

    st.markdown("""
    <div class="info-card" style="margin-top:4px">
      <div style="font-size:12px;color:#555;line-height:2">
        <span style="color:#F5C518;font-weight:700">How it works:</span><br>
        1. Upload your Amazon invoice batch PDF<br>
        2. Click <b style="color:#F5C518">⚡ PROCESS INVOICES</b><br>
        3. Download the stamped output PDF<br><br>
        <span style="color:#F5C518;font-weight:700">Box format on every invoice:</span><br>
        ████████████████████████████<br>
        █  PRODUCT CODE  |  QTY: X  █<br>
        █    ★ ★ ★  (if qty &gt; 1)    █<br>
        ████████████████████████████
      </div>
    </div>
    """, unsafe_allow_html=True)
